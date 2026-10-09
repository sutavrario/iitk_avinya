"""Generate the ingestion sample files (all data is fictional).

Run from the repo root:
    backend/.venv/bin/python sample_data/ingestion/generate_samples.py
Requires the backend dev dependencies (openpyxl, xlwt, fpdf2, Pillow).
"""

from datetime import datetime
from pathlib import Path

import openpyxl
import xlwt
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = Path(__file__).parent


def sales_register_xlsx() -> None:
    """Clean register with standard headers; CGST+SGST columns must be summed as tax."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Invoice No", "Customer Name", "Invoice Date", "Due Date", "Taxable Value", "CGST", "SGST", "Invoice Value", "Status", "UTR"])
    rows = [
        ("INV-1001", "Patel Agencies", datetime(2026, 9, 1), datetime(2026, 9, 30), 41100.0, 3699.0, 3699.0, 48498.0, "Unpaid", None),
        ("INV-1002", "Gupta Kirana", datetime(2026, 9, 3), datetime(2026, 10, 3), 12000.0, 300.0, 300.0, 12600.0, "Paid", "UTR123456789"),
        ("INV-1003", "Mehta Stores", datetime(2026, 9, 10), datetime(2026, 10, 10), 5000.0, 450.0, 450.0, 5900.0, None, None),
        ("INV-1004", "Sharma Traders", datetime(2026, 9, 15), datetime(2026, 10, 15), 20000.0, 1800.0, 1800.0, 23600.0, "Pending", None),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(OUT / "sales_register.xlsx")


def tally_export_xlsx() -> None:
    """Tally-style export: title rows above the header, odd column names, a totals row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Register"
    ws.append(["Sharma General Store"])
    ws.append(["Sales Register"])
    ws.append(["1-Apr-2026 to 30-Sep-2026"])
    ws.append([])
    ws.append(["Date", "Particulars", "Voucher No.", "Taxable Amt", "IGST", "Gross Total"])
    ws.append([datetime(2026, 8, 2), "Kumar Electricals", "S/245", 10000, 1800, 11800])
    ws.append([datetime(2026, 8, 9), "Rao & Sons", "S/246", 25000, 4500, 29500])
    ws.append([datetime(2026, 8, 21), "Kumar Electricals", "S/247", 8000, 1440, 9440])
    ws.append([None, "Grand Total", None, 43000, 7740, 50740])
    wb.create_sheet("Notes").append(["Generated for VyaparAI tests"])
    wb.save(OUT / "tally_sales_export.xlsx")


def legacy_xls() -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Bills")
    for c, h in enumerate(["Bill No", "Party Name", "Bill Date", "Amount"]):
        ws.write(0, c, h)
    data = [("B-501", "Joshi Textiles", "05/08/2026", "15,000.00"), ("B-502", "Iyer Medicals", "18/08/2026", "8,250.50")]
    for r, row in enumerate(data, start=1):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    wb.save(str(OUT / "legacy_bills.xls"))


def problems_csv() -> None:
    """Every row exercises a validation rule. See README.md for what each row tests."""
    text = """Inv #,Buyer,Date,Due,Sub Total,GST,Total,Payment Status
INV-2001,Patel Agencies,01/09/2026,30/09/2026,"₹10,000.00","₹1,800.00","₹11,800.00",Unpaid
INV-2002,Gupta Kirana,02/09/2026,02/10/2026,5000,900,,Due
INV-2003,Mehta Stores,03/09/2026,03/10/2026,8000,1440,9999,
INV-2001,Patel Agencies,01/09/2026,30/09/2026,10000,1800,11800,Unpaid
INV-2004,,04/09/2026,04/10/2026,2000,360,2360,
INV-2005,Sharma Traders,31/02/2026,31/03/2026,3000,540,3540,
INV-2006,Rao & Sons,15/09/2026,10/09/2026,"Rs. 500/-",90,590,PAID
,Kumar Electricals,16/09/2026,16/10/2026,1000,180,1180,
"""
    (OUT / "sales_with_problems.csv").write_text(text, encoding="utf-8")


def purchase_csv() -> None:
    text = """Supplier;Bill No;Bill Date;Taxable Value;Total GST;Amount Payable;Payment Ref
Bharat Wholesale Supplies;BWS/88/2026;2026-09-05;50000;9000;59000;NEFT N12345678
City Power Distribution;CPD-7731;2026-09-12;4200;756;4956;
"""
    (OUT / "purchase_bills.csv").write_text(text, encoding="utf-8")


def unmapped_csv() -> None:
    """Headers give no clue — must go to the column-mapping step."""
    text = """Col1,Col2,Col3,Col4
A-1,Patel Agencies,2026-09-01,1180
A-2,Gupta Kirana,2026-09-02,590
"""
    (OUT / "unknown_columns.csv").write_text(text, encoding="utf-8")


def duplicate_csv() -> None:
    """Repeats INV-1001 from sales_register.xlsx to test cross-document duplicate detection."""
    text = """Invoice No,Customer Name,Invoice Date,Due Date,Taxable Value,Total Tax,Invoice Value
INV-1001,Patel Agencies,2026-09-01,2026-09-30,41100,7398,48498
INV-1099,New Customer Pvt Ltd,2026-09-20,2026-10-20,1000,180,1180
"""
    (OUT / "duplicate_of_register.csv").write_text(text, encoding="utf-8")


INVOICE_LINES = [
    "Sharma General Store",
    "12 MG Road, Pune 411001  GSTIN: 27ABCDE1234F1Z5",
    "TAX INVOICE",
    "Invoice No: SGS/2026/0457",
    "Invoice Date: 18/09/2026",
    "Due Date: 18/10/2026",
    "Bill To: Deshmukh Caterers",
    "Item                     Qty     Rate      Amount",
    "Basmati Rice 25kg          4  2,150.00   8,600.00",
    "Sunflower Oil 15L          2  2,390.00   4,780.00",
    "Sub Total: Rs. 13,380.00",
    "CGST @ 2.5%: Rs. 334.50",
    "SGST @ 2.5%: Rs. 334.50",
    "Grand Total: Rs. 14,049.00",
    "Payment Status: Unpaid",
    "Thank you for your business!",
]


def digital_pdf() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in INVOICE_LINES:
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(OUT / "invoice_digital.pdf"))

    # Encrypted copy to test the password-protected path.
    enc = FPDF()
    enc.set_encryption(owner_password="owner-secret", user_password="open-me")
    enc.add_page()
    enc.set_font("Helvetica", size=11)
    enc.cell(0, 8, "Invoice No: LOCKED-1")
    enc.output(str(OUT / "malformed" / "password_protected.pdf"))


def _render(lines: list[str], size: int = 34, rotate: float = 0.0, blur: bool = False) -> Image.Image:
    font = ImageFont.load_default(size=size)
    img = Image.new("RGB", (1700, 80 + len(lines) * int(size * 1.6)), "white")
    draw = ImageDraw.Draw(img)
    y = 40
    for line in lines:
        draw.text((60, y), line, fill="black", font=font)
        y += int(size * 1.6)
    if rotate:
        img = img.rotate(rotate, expand=True, fillcolor="white")
    if blur:
        img = img.filter(ImageFilter.GaussianBlur(0.6))
    return img


def scanned_pdf_and_photo() -> None:
    _render(INVOICE_LINES).save(OUT / "invoice_scanned.pdf", resolution=150)  # image-only PDF
    photo_lines = [
        "Bharat Wholesale Supplies",
        "Invoice No: BWS/91/2026",
        "Invoice Date: 22/09/2026",
        "Bill To: Sharma General Store",
        "Taxable Value: Rs. 20,000.00",
        "IGST @ 18%: Rs. 3,600.00",
        "Grand Total: Rs. 23,600.00",
    ]
    _render(photo_lines, rotate=0.8, blur=True).save(OUT / "bill_photo.png")


def malformed() -> None:
    m = OUT / "malformed"
    m.mkdir(exist_ok=True)
    (m / "corrupt.xlsx").write_bytes(b"PK\x03\x04" + b"this is not a real workbook" * 20)
    full = (OUT / "invoice_digital.pdf").read_bytes()
    (m / "truncated.pdf").write_bytes(full[:300])
    (m / "header_only.csv").write_text("Invoice No,Customer,Date,Total\n", encoding="utf-8")
    (m / "ragged.csv").write_text('Invoice No,Customer,Total\nINV-1,"Unclosed quote,100\nINV-2,B,200,extra,cols\n', encoding="utf-8")


if __name__ == "__main__":
    (OUT / "malformed").mkdir(exist_ok=True)
    sales_register_xlsx()
    tally_export_xlsx()
    legacy_xls()
    problems_csv()
    purchase_csv()
    unmapped_csv()
    duplicate_csv()
    digital_pdf()
    scanned_pdf_and_photo()
    malformed()
    print("Samples written to", OUT)
